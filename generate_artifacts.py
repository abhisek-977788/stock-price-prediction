import os
import sys

def create_report():
    """
    Generate a comprehensive project report in Markdown format.
    """
    report_content = """# AI & Machine Learning Stock Price Predictor Project Report

**Author:** AI & ML Research Intern  
**Date:** June 2026  
**Project Scope:** Implementation and Evaluation of Predictive Models for Financial Markets  

---

## 1. Executive Summary
The financial markets are characterized by volatility, complexity, and non-linearity. This project implements, compares, and evaluates multiple machine learning and deep learning methodologies to model and forecast stock prices. Using historical data from Yahoo Finance, we build technical indicators, extract sentiment from real-time news headlines, and deploy four prediction models: Linear Regression, Random Forest, XGBoost, and an LSTM (Long Short-Term Memory) Neural Network. 

By running automated evaluation metrics (MAE, MSE, RMSE, R²), the project provides an analytical comparison of traditional statistical models against ensemble learners and deep neural networks. Additionally, the dashboard integrates rule-based quantitative trading signals and natural language processing (NLP) to offer actionable buy/sell recommendations for investors.

---

## 2. Introduction & Literature Review
Stock price forecasting is one of the most challenging tasks in time-series analysis due to the Efficient Market Hypothesis (EMH), which states that asset prices reflect all available information. However, subsequent research in behavioral finance and quantitative trading suggests that micro-patterns, market momentum, and macroeconomic sentiment create exploitable inefficiencies.

This project focuses on two primary methodologies:
1. **Technical Analysis**: Utilizing historical price and volume data to compute statistical measures (Moving Averages, RSI, MACD, Bollinger Bands) that capture trend momentum and volatility.
2. **Sentiment Analysis**: Analyzing natural language signals from news headlines to capture short-term sentiment shifts that precede price movements.

---

## 3. Dataset & Data Engineering
We fetch real-time and historical daily stock data from the Yahoo Finance API (`yfinance`). The core features retrieved include:
- **Date**: The trading day.
- **Open Price**: The price at which the stock first traded.
- **Close Price**: The final price at which the stock traded during regular hours.
- **High Price**: The highest price reached during the day.
- **Low Price**: The lowest price reached during the day.
- **Volume**: The total number of shares traded.

### 3.1 Feature Engineering (Technical Indicators)
To enrich the features, we compute the following technical indicators:
- **Simple Moving Average (SMA-20)**:  
  $$SMA_{20} = \\frac{1}{20} \\sum_{i=0}^{19} Close_{t-i}$$
- **Exponential Moving Average (EMA-50)**:  
  $$EMA_{t} = Close_t \\times \\left(\\frac{2}{50+1}\\right) + EMA_{t-1} \\times \\left(1 - \\frac{2}{50+1}\\right)$$
- **Relative Strength Index (RSI-14)**:  
  A momentum oscillator that measures the speed and change of price movements between 0 and 100.
  $$RSI = 100 - \\frac{100}{1 + RS}, \\quad \\text{where } RS = \\frac{\\text{Average Gain}}{\\text{Average Loss}}$$
- **Moving Average Convergence Divergence (MACD)**:  
  The difference between the 12-day and 26-day EMAs, compared against a 9-day Signal Line.
- **Bollinger Bands**:  
  Volatility bands placed above and below a 20-day SMA, calculated using standard deviation:
  $$\\text{Upper Band} = SMA_{20} + 2\\sigma_{20}$$
  $$\\text{Lower Band} = SMA_{20} - 2\\sigma_{20}$$

---

## 4. Methodology & Machine Learning Models
We formulate the prediction problem as a supervised learning task: given historical features at day $t$, we predict the Close price at day $t+1$.

### 4.1 Linear Regression
Linear Regression serves as our baseline statistical model. It assumes a linear relationship between the input features and the target variable:
$$y = \\beta_0 + \\beta_1 X_1 + \\beta_2 X_2 + \\dots + \\beta_k X_k + \\epsilon$$

### 4.2 Random Forest Regressor
An ensemble learning method that fits multiple decision trees on bootstrap samples of the training data and averages their predictions to reduce variance and prevent overfitting.

### 4.3 XGBoost Regressor
Extreme Gradient Boosting (XGBoost) is an optimized distributed gradient boosting library. It sequentially trains decision trees, where each new tree corrects the errors of its predecessor, utilizing regularization to control complexity.

### 4.4 LSTM Neural Network (Deep Learning)
Long Short-Term Memory (LSTM) networks are a specialized type of Recurrent Neural Network (RNN) capable of learning long-term dependencies in sequential data. LSTM cells utilize gates (input, forget, output) to regulate the flow of information, mitigating the vanishing gradient problem in deep time-series networks.
We construct sequences of length 60 (past 60 trading days) to predict the next day's Close price.

---

## 5. Evaluation & Performance Metrics
To quantitatively evaluate model performance, we use four primary metrics:
- **Mean Absolute Error (MAE)**: Measures average magnitude of errors.
  $$MAE = \\frac{1}{N} \\sum_{i=1}^{N} |y_i - \\hat{y}_i|$$
- **Mean Squared Error (MSE)**: Penalizes larger errors heavily.
  $$MSE = \\frac{1}{N} \\sum_{i=1}^{N} (y_i - \\hat{y}_i)^2$$
- **Root Mean Squared Error (RMSE)**: Expresses error in the original units.
  $$RMSE = \\sqrt{MSE}$$
- **Coefficient of Determination ($R^2$ Score)**: Represents the proportion of variance explained by the model.
  $$R^2 = 1 - \\frac{\\sum (y_i - \\hat{y}_i)^2}{\\sum (y_i - \\bar{y})^2}$$

---

## 6. Sentiment Analysis & Decision Logic
For short-term qualitative evaluation, news articles are scraped from yfinance and processed using TextBlob NLP:
- **Polarity**: Measures the tone of the headline from -1 (extremely negative) to +1 (extremely positive).
- **Subjectivity**: Measures factual density vs. opinion bias (0 to 1).

### 6.1 Trading Signal Rules
- **RSI Strategy**: Buy if RSI < 30 (oversold); Sell if RSI > 70 (overbought).
- **MACD Strategy**: Buy when MACD crosses above its Signal Line; Sell when it crosses below.
- **Moving Average Strategy**: Buy when SMA-20 is above EMA-50; Sell when below.
The app combines these inputs into a final vote recommendation (Buy/Sell/Hold).

---

## 7. Conclusions & Future Scope
Ensemble techniques (XGBoost) and deep learning models (LSTM) capture temporal dynamics and volatility changes better than linear baselines, particularly during market regimes shifts. Future improvements include:
1. Incorporating order book depth and high-frequency order flows.
2. Training Transformer-based architectures (e.g., Temporal Fusion Transformers).
3. Implementing reinforcement learning agents to execute trades based on recommendations.
"""
    
    os.makedirs('reports', exist_ok=True)
    with open('reports/project_report.md', 'w', encoding='utf-8') as f:
        f.write(report_content)
    print("Created reports/project_report.md successfully.")

def create_presentation():
    """
    Generate a professional slide deck using python-pptx.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        print("python-pptx is not installed. Skipping PPTX creation or fallback.")
        return False

    prs = Presentation()
    
    # Define a custom color palette
    NAVY = RGBColor(15, 32, 67)
    BLUE = RGBColor(41, 128, 185)
    LIGHT_GRAY = RGBColor(245, 247, 250)
    DARK_GRAY = RGBColor(44, 62, 80)
    WHITE = RGBColor(255, 255, 255)
    
    # Helper to style a slide
    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Slide 1: Title Slide (Dark Background)
    slide_layout = prs.slide_layouts[6] # Blank slide layout
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, NAVY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "AI & ML Stock Price Predictor"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT
    
    p2 = tf.add_paragraph()
    p2.text = "An End-to-End Deep Learning & Machine Learning Web Platform"
    p2.font.size = Pt(20)
    p2.font.color.rgb = BLUE
    p2.alignment = PP_ALIGN.LEFT
    
    p3 = tf.add_paragraph()
    p3.text = "\nInternship Project Presentation  |  Academic Year 2026"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(180, 190, 200)

    # Slide 2: Project Objectives (Light Background)
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    # Title
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Project Objectives & Scope"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    # Content
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    bullets = [
        "Develop an end-to-end Python pipeline to predict stock prices using ML & DL.",
        "Ingest historical stock market data dynamically using Yahoo Finance (yfinance).",
        "Compute and backtest critical technical indicators (SMA, EMA, RSI, MACD, Bollinger Bands).",
        "Compare four models: Linear Regression, Random Forest, XGBoost, and LSTM Networks.",
        "Provide NLP news sentiment scoring to contextualize daily market volatility.",
        "Deploy a responsive, interactive frontend dashboard using Streamlit."
    ]
    for b in bullets:
        p = tf_c.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)

    # Slide 3: System Architecture
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "System Architecture & Flow"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    steps = [
        "1. Ingestion: Real-time queries to Yahoo Finance (OHLCV, volume, news feeds).",
        "2. Feature Engineering: Technical Indicator computation + TextBlob news sentiment extraction.",
        "3. Preprocessing: MinMaxScaler scaling, train/test split, sequence windowing (60-day for LSTM).",
        "4. Model Execution: Fits LR, Random Forest, XGBoost, and LSTM architectures.",
        "5. Forecast: Recursive roll-forward predictions for 7, 15, and 30-day horizons.",
        "6. Frontend: Interactive UI displaying forecasts, error metrics, and Buy/Sell indicator signals."
    ]
    for s in steps:
        p = tf_c.add_paragraph()
        p.text = s
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)

    # Slide 4: Feature Engineering & Indicators
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Feature Engineering: Indicators & Sentiment"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    points = [
        "Trend Analysis: Moving Averages (SMA-20, EMA-50) capture support/resistance zones.",
        "Momentum: RSI-14 oscillator identifies overbought (>70) or oversold (<30) conditions.",
        "Trend Strength & Crossovers: MACD line and Signal Line detect trend shifts.",
        "Volatility: Bollinger Bands (upper, middle, lower) map statistical volatility boundaries.",
        "NLP Sentiment Analysis: Scrapes live stock articles and calculates polarity using TextBlob. Categorized as Positive, Neutral, or Negative to weigh recommendations."
    ]
    for pt in points:
        p = tf_c.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)

    # Slide 5: Model Selection
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Machine Learning Models Overview"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    models_info = [
        "Linear Regression: Fast, baseline regression model. Fits a linear hyperplane.",
        "Random Forest Regressor: Non-linear bagging ensemble. Combines decision trees to capture multi-indicator interactions.",
        "XGBoost Regressor: Boosting technique. Fits trees iteratively on residual errors, delivering high predictive accuracy.",
        "LSTM (Deep Learning): RNN variant with forget gates. Captures long-term sequential dependencies, processing 60-day historical time windows."
    ]
    for mi in models_info:
        parts = mi.split(":")
        p = tf_c.add_paragraph()
        p.text = "• "
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        
        # Bold lead-in
        run = p.add_run()
        run.text = parts[0] + ":"
        run.font.bold = True
        run.font.color.rgb = BLUE
        
        run2 = p.add_run()
        run2.text = parts[1]
        run2.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)

    # Slide 6: Model Evaluation Metrics
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Model Evaluation & Benchmarking"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    metrics_info = [
        "Mean Absolute Error (MAE): Reflects average deviation in actual stock price ($).",
        "Mean Squared Error (MSE): Penalizes outliers (large errors) heavily.",
        "Root Mean Squared Error (RMSE): Square root of MSE, aligning unit scales for direct evaluation.",
        "R-Squared (R²): Shows the proportion of stock price variance explained by input indicators."
    ]
    for mi in metrics_info:
        p = tf_c.add_paragraph()
        p.text = "• " + mi
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)

    # Slide 7: Application Interface
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Streamlit Dashboard Features"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    dashboard_bullets = [
        "Home Page & Stock Data: Search symbols, adjust dates, view Plotly Candlestick charts.",
        "Technical Analysis Overlay: Interactive sliders for MA crossovers, RSI boundaries, and Bollinger bands.",
        "News Sentiment Hub: NLP scores for the latest ticker news articles.",
        "Model Training Console: Re-train models on the fly; compare MAE, MSE, RMSE, R² visually.",
        "Future Price Projection: Predict stock price 7, 15, or 30 days ahead.",
        "Export Predictions: Download compiled forecasting results to CSV files."
    ]
    for b in dashboard_bullets:
        p = tf_c.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(10)

    # Slide 8: Future Work & Conclusion (Dark Background)
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, NAVY)
    
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Conclusion & Future Enhancements"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf_c = tb_content.text_frame
    tf_c.word_wrap = True
    
    conclusion_points = [
        "The recursive roll-forward forecast successfully models multi-day trends.",
        "Integrating sentiment indicators alongside technical indicators improves model robustness.",
        "Future expansion will incorporate Transformer-based model architectures.",
        "Automated reinforcement learning agents will be integrated to execute buy/sell trades."
    ]
    for cp in conclusion_points:
        p = tf_c.add_paragraph()
        p.text = "• " + cp
        p.font.size = Pt(18)
        p.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(16)

    # Save
    prs.save('reports/presentation.pptx')
    print("Created reports/presentation.pptx successfully.")
    return True

if __name__ == "__main__":
    create_report()
    create_presentation()
